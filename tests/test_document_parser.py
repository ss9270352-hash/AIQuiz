"""
test_document_parser.py - Tests for TXT / PDF / PPTX upload extraction.

Covers:
1. Plain-text decoding and metadata.
2. PPTX slide text, tables, and speaker notes extraction.
3. PDF page text extraction.
4. Honest errors for empty, corrupted, unsupported, and text-less files
   (e.g. scanned PDFs with no text layer).
"""

import io

import pytest

from document_parser import (
    DocumentParserError,
    extract_text_from_upload,
)


# -----------------------------------------------------------------------------
# Helpers: build real .pptx / .pdf files in memory
# -----------------------------------------------------------------------------

def _make_pptx_bytes() -> bytes:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide1.shapes.title.text = "Introduction to Operating Systems"
    slide1.placeholders[1].text = "Processes and Threads\nVirtual Memory"

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Algorithm"
    table.cell(0, 1).text = "Complexity"
    table.cell(1, 0).text = "Merge Sort"
    table.cell(1, 1).text = "O(n log n)"
    slide2.notes_slide.notes_text_frame.text = "Remember: scheduling is on the exam."

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _make_pdf_bytes(lines) -> bytes:
    """Builds a minimal, valid one-page text PDF with correct xref offsets."""
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    text_ops = "\n".join(
        f"BT /F1 12 Tf 50 {750 - i * 20} Td ({escape(line)}) Tj ET"
        for i, line in enumerate(lines)
    )
    content = text_ops.encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_start = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += f"{offset:010d} 00000 n \n".encode()
    pdf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_start}\n%%EOF\n"
    ).encode()
    return bytes(pdf)


def _make_blank_pdf_bytes() -> bytes:
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# -----------------------------------------------------------------------------
# TXT
# -----------------------------------------------------------------------------

def test_txt_extraction():
    parsed = extract_text_from_upload(
        "Lecture notes: a stack is LIFO, a queue is FIFO.".encode("utf-8"),
        "notes.txt",
    )
    assert parsed.file_type == "txt"
    assert parsed.unit_count == 1
    assert "stack is LIFO" in parsed.text
    assert parsed.char_count == len(parsed.text.strip())


def test_txt_uppercase_extension():
    parsed = extract_text_from_upload(b"Some lecture notes content.", "NOTES.TXT")
    assert parsed.file_type == "txt"


# -----------------------------------------------------------------------------
# PPTX
# -----------------------------------------------------------------------------

def test_pptx_extraction_text_tables_and_notes():
    parsed = extract_text_from_upload(_make_pptx_bytes(), "lecture.pptx")
    assert parsed.file_type == "pptx"
    assert parsed.unit_count == 2
    assert parsed.unit_label == "slide"
    assert "Introduction to Operating Systems" in parsed.text
    assert "Processes and Threads" in parsed.text
    assert "--- Slide 1 ---" in parsed.text
    assert "--- Slide 2 ---" in parsed.text
    # Table cell text is captured
    assert "Merge Sort" in parsed.text
    assert "O(n log n)" in parsed.text
    # Speaker notes are captured
    assert "scheduling is on the exam" in parsed.text


def test_pptx_corrupted_file_error():
    with pytest.raises(DocumentParserError):
        extract_text_from_upload(b"this is not a powerpoint file", "slides.pptx")


def test_old_ppt_extension_gets_specific_error():
    with pytest.raises(DocumentParserError, match="\\.pptx"):
        extract_text_from_upload(b"anything", "old_slides.ppt")


# -----------------------------------------------------------------------------
# PDF
# -----------------------------------------------------------------------------

def test_pdf_extraction():
    pdf_bytes = _make_pdf_bytes([
        "Operating systems schedule processes.",
        "Virtual memory uses paging.",
    ])
    parsed = extract_text_from_upload(pdf_bytes, "chapter.pdf")
    assert parsed.file_type == "pdf"
    assert parsed.unit_count == 1
    assert parsed.unit_label == "page"
    assert "--- Page 1 ---" in parsed.text
    assert "Operating systems schedule processes." in parsed.text
    assert "Virtual memory uses paging." in parsed.text


def test_pdf_without_text_layer_error():
    # A blank / scanned-style PDF has no extractable text: honest error, no crash.
    with pytest.raises(DocumentParserError, match="No readable text"):
        extract_text_from_upload(_make_blank_pdf_bytes(), "scanned.pdf")


def test_pdf_corrupted_file_error():
    with pytest.raises(DocumentParserError):
        extract_text_from_upload(b"%PDF- not really a pdf at all", "broken.pdf")


# -----------------------------------------------------------------------------
# Dispatcher / shared cases
# -----------------------------------------------------------------------------

def test_empty_file_error():
    with pytest.raises(DocumentParserError, match="empty"):
        extract_text_from_upload(b"", "notes.txt")


def test_unsupported_extension_error():
    with pytest.raises(DocumentParserError, match="Unsupported file type"):
        extract_text_from_upload(b"some content", "notes.docx")
