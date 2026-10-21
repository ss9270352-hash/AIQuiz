"""
document_parser.py - Text extraction from uploaded study documents.

Supports the file types a student actually has lecture material in:
  - .txt   plain text notes (UTF-8 / UTF-16 / Latin-1 tolerant)
  - .pdf   text-based PDFs, extracted page by page with pypdf
  - .pptx  PowerPoint decks, extracted slide by slide with python-pptx
           (text frames, tables, grouped shapes, and speaker notes)

Every extractor returns a ParsedDocument with the plain text plus
metadata (type, page/slide count) so the UI can tell the student
exactly what was loaded. Failures raise DocumentParserError with a
plain-language message that is safe to show directly in the UI.

Known limits (by design, and reported honestly to the user):
  - Scanned / image-only PDFs contain no text layer, so nothing can be
    extracted from them without OCR. That case raises a clear error.
  - Images, charts and diagrams inside PDFs/PPTX files are skipped;
    only their captions / surrounding text are captured.
"""

import io
import os
from dataclasses import dataclass
from typing import List, Optional


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".pptx"}


class DocumentParserError(Exception):
    """User-friendly error indicating a document could not be parsed."""
    pass


@dataclass
class ParsedDocument:
    """Result of extracting text from one uploaded file."""
    text: str
    file_name: str
    file_type: str          # "txt" | "pdf" | "pptx"
    unit_count: int         # pages for PDF, slides for PPTX, 1 for TXT
    unit_label: str         # "page" | "slide" | "file"

    @property
    def char_count(self) -> int:
        return len(self.text.strip())


# -----------------------------------------------------------------------------
# TXT
# -----------------------------------------------------------------------------

def extract_text_from_txt(data: bytes, file_name: str = "notes.txt") -> ParsedDocument:
    """Decodes plain-text bytes, tolerating common student-file encodings."""
    if not data:
        raise DocumentParserError("That text file is empty — there are no notes to build a quiz from.")

    text: Optional[str] = None
    # utf-8-sig strips a Windows BOM; utf-16 covers Notepad's "Unicode" saves.
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        # Last resort: Latin-1 never fails and keeps the words readable.
        text = data.decode("latin-1", errors="replace")

    if not text.strip():
        raise DocumentParserError("That text file contains no readable text.")

    return ParsedDocument(
        text=text.strip(),
        file_name=file_name,
        file_type="txt",
        unit_count=1,
        unit_label="file",
    )


# -----------------------------------------------------------------------------
# PDF (pypdf)
# -----------------------------------------------------------------------------

def extract_text_from_pdf(data: bytes, file_name: str = "notes.pdf") -> ParsedDocument:
    """Extracts the text layer of a PDF, page by page."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocumentParserError(
            "PDF support is not installed. Run 'pip install -r requirements.txt' and restart the app."
        )

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise DocumentParserError(
            "That PDF could not be opened — the file may be corrupted or not a real PDF."
        )

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise DocumentParserError(
                "That PDF is password-protected. Remove the password and upload it again."
            )

    page_texts: List[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            page_texts.append(f"--- Page {index} ---\n{page_text.strip()}")

    if not page_texts:
        raise DocumentParserError(
            "No readable text was found in that PDF. If it is a scanned document or "
            "photos of pages, it has no text layer to extract — please paste the notes "
            "as text instead, or upload a text-based PDF."
        )

    return ParsedDocument(
        text="\n\n".join(page_texts),
        file_name=file_name,
        file_type="pdf",
        unit_count=len(reader.pages),
        unit_label="page",
    )


# -----------------------------------------------------------------------------
# PPTX (python-pptx)
# -----------------------------------------------------------------------------

def _collect_shape_texts(shape, texts: List[str]) -> None:
    """Recursively collects text from one shape: text frames, tables, groups."""
    # Grouped shapes contain child shapes — walk into them.
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            _collect_shape_texts(child, texts)

    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_texts:
                texts.append(" | ".join(cell_texts))
        return

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            line = "".join(run.text for run in paragraph.runs).strip()
            if not line and paragraph.text.strip():
                line = paragraph.text.strip()
            if line:
                texts.append(line)


def extract_text_from_pptx(data: bytes, file_name: str = "slides.pptx") -> ParsedDocument:
    """Extracts slide text (and speaker notes) from a PowerPoint deck."""
    try:
        from pptx import Presentation
    except ImportError:
        raise DocumentParserError(
            "PowerPoint support is not installed. Run 'pip install -r requirements.txt' and restart the app."
        )

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception:
        raise DocumentParserError(
            "That PowerPoint file could not be opened — it may be corrupted, "
            "password-protected, or an older .ppt file renamed to .pptx."
        )

    slide_texts: List[str] = []
    slide_count = 0
    for index, slide in enumerate(presentation.slides, start=1):
        slide_count = index
        texts: List[str] = []
        for shape in slide.shapes:
            _collect_shape_texts(shape, texts)

        # Speaker notes often hold the real lecture content — include them.
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                texts.append(f"[Speaker notes] {notes_text}")

        if texts:
            slide_texts.append(f"--- Slide {index} ---\n" + "\n".join(texts))

    if not slide_texts:
        raise DocumentParserError(
            "No readable text was found in that presentation. If the slides are "
            "only images or screenshots, please paste the notes as text instead."
        )

    return ParsedDocument(
        text="\n\n".join(slide_texts),
        file_name=file_name,
        file_type="pptx",
        unit_count=slide_count,
        unit_label="slide",
    )


# -----------------------------------------------------------------------------
# Dispatcher
# -----------------------------------------------------------------------------

def extract_text_from_upload(data: bytes, file_name: str) -> ParsedDocument:
    """
    Main entry point: routes an uploaded file to the right extractor
    based on its extension.
    """
    if data is None or len(data) == 0:
        raise DocumentParserError("That file is empty — there are no notes to build a quiz from.")

    ext = os.path.splitext(file_name or "")[1].lower()
    if ext == ".txt":
        return extract_text_from_txt(data, file_name)
    if ext == ".pdf":
        return extract_text_from_pdf(data, file_name)
    if ext == ".pptx":
        return extract_text_from_pptx(data, file_name)
    if ext == ".ppt":
        raise DocumentParserError(
            "Old .ppt files are not supported. Open the deck in PowerPoint, "
            "save it as .pptx, and upload that instead."
        )
    raise DocumentParserError(
        f"Unsupported file type '{ext or 'unknown'}'. Please upload a .txt, .pdf, or .pptx file."
    )
